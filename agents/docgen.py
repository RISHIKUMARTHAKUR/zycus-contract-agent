"""
Document Generation Tool.

Deterministic tool call: takes the drafted sections + flags and produces a
.docx. No LLM involved here -- this is pure formatting, treated as a callable
tool the agent pipeline invokes once drafting is complete.
"""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from agents.validator import Flag

HIGHLIGHT_YELLOW = RGBColor(0x8A, 0x6D, 0x00)


def generate_docx(sections: list[dict], flags: list[Flag], output_path: str) -> str:
    doc = Document()

    title = doc.add_heading("MUTUAL NON-DISCLOSURE AGREEMENT", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for section in sections:
        if section["heading"]:
            doc.add_heading(section["heading"], level=2)
        para = doc.add_paragraph(section["body"])
        para.paragraph_format.space_after = Pt(10)

    review_flags = [f for f in flags if f.resolution == "needs_human_review"]
    if review_flags:
        doc.add_page_break()
        doc.add_heading("Items Flagged for Human Review", level=1)
        doc.add_paragraph(
            "The clauses below were not finalized automatically because they "
            "involve a change to one party's rights or exposure. Please "
            "review and approve before sending."
        )
        for f in review_flags:
            p = doc.add_paragraph()
            run = p.add_run(f"[{f.field}] {f.issue}")
            run.bold = True
            run.font.color.rgb = HIGHLIGHT_YELLOW
            if f.detail:
                doc.add_paragraph(f.detail)

    doc.save(output_path)
    return output_path
